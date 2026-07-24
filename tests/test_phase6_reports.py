from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook
from pypdf import PdfReader


def test_admin_can_open_report_center(client, login, admin_user, phase4_assessment):
    login(admin_user.email)
    response = client.get(f"/reports/assessments/{phase4_assessment.public_id}")
    assert response.status_code == 200
    assert b"Centro de reportes" in response.data


def test_admin_can_export_excel(client, login, admin_user, phase4_assessment):
    login(admin_user.email)
    response = client.get(f"/reports/assessments/{phase4_assessment.public_id}/excel")
    assert response.status_code == 200
    assert response.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    workbook = load_workbook(BytesIO(response.data), read_only=True)
    assert "Resumen ejecutivo" in workbook.sheetnames
    assert "Preguntas" in workbook.sheetnames
    assert "Bitácora" in workbook.sheetnames


def test_admin_can_export_pdf(client, login, admin_user, phase4_assessment):
    login(admin_user.email)
    response = client.get(f"/reports/assessments/{phase4_assessment.public_id}/pdf")
    assert response.status_code == 200
    assert response.mimetype == "application/pdf"
    assert response.data.startswith(b"%PDF-")
    reader = PdfReader(BytesIO(response.data))
    assert len(reader.pages) >= 5
    assert reader.metadata.title.startswith("Assessment DevSecOps")


def test_respondent_cannot_export_unpublished_results(
    client, login, respondent_user, phase4_assessment
):
    login(respondent_user.email)
    response = client.get(f"/reports/assessments/{phase4_assessment.public_id}/pdf")
    assert response.status_code == 403
