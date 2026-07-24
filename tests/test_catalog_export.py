from pathlib import Path

from openpyxl import load_workbook

from app.services.catalog_export_service import catalog_export_service
from app.services.samm_import_service import catalog_import_service

SOURCE = Path(__file__).parents[1] / "data" / "SAMM_spreadsheet.xlsx"


def test_export_generates_compatible_workbook(app, admin_user):
    record = catalog_import_service.create_preview_from_path(SOURCE, actor_id=admin_user.id)
    version = catalog_import_service.confirm(
        record,
        version_name="Export",
        version_number="export-1",
        description=None,
        publish=False,
        actor_id=admin_user.id,
    )
    output = catalog_export_service.build_workbook(version)
    workbook = load_workbook(output, read_only=True, data_only=True)
    assert workbook.sheetnames == ["Metadata", "imp-questions", "imp-answers"]
    assert workbook["imp-questions"].max_row == 91
    assert workbook["imp-answers"].max_row == 25
    workbook.close()
