from __future__ import annotations

import hashlib
import json
import re
import zipfile
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

QUESTION_HEADERS = (
    "ID",
    "Business Function",
    "Security Practice",
    "Activity",
    "Maturity",
    "Question",
    "Guidance",
    "Answer Option",
)
ANSWER_HEADERS = ("ANS_SET_CODE", "A", "B", "C", "D", "A_W", "B_W", "C_W", "D_W")
QUESTION_ID_PATTERN = re.compile(r"^[A-Za-z0-9]+(?:-[A-Za-z0-9]+){4,}$")


@dataclass(slots=True)
class ImportIssue:
    sheet: str
    row: int | None
    field: str | None
    message: str
    severity: str = "error"

    def as_dict(self) -> dict[str, Any]:
        return {
            "sheet": self.sheet,
            "row": self.row,
            "field": self.field,
            "message": self.message,
            "severity": self.severity,
        }


@dataclass(slots=True)
class ImportPreview:
    source_name: str
    source_file_hash: str
    source_version: str | None
    answer_sets: list[dict[str, Any]] = field(default_factory=list)
    questions: list[dict[str, Any]] = field(default_factory=list)
    issues: list[ImportIssue] = field(default_factory=list)

    @property
    def errors(self) -> list[ImportIssue]:
        return [issue for issue in self.issues if issue.severity == "error"]

    @property
    def warnings(self) -> list[ImportIssue]:
        return [issue for issue in self.issues if issue.severity == "warning"]

    @property
    def summary(self) -> dict[str, int]:
        questions = self.questions
        return {
            "business_functions": len({item["business_function_code"] for item in questions}),
            "security_practices": len({item["security_practice_code"] for item in questions}),
            "practice_streams": len({item["practice_stream_code"] for item in questions}),
            "maturity_levels": len({item["maturity_level"] for item in questions}),
            "questions": len(questions),
            "answer_sets": len(self.answer_sets),
            "quality_criteria": sum(len(item["criteria"]) for item in questions),
            "errors": len(self.errors),
            "warnings": len(self.warnings),
        }

    def as_dict(self) -> dict[str, Any]:
        return {
            "source_name": self.source_name,
            "source_file_hash": self.source_file_hash,
            "source_version": self.source_version,
            "summary": self.summary,
            "answer_sets": self.answer_sets,
            "questions": self.questions,
            "issues": [issue.as_dict() for issue in self.issues],
        }


class SammWorkbookParser:
    def parse(self, file_path: str | Path) -> ImportPreview:
        path = Path(file_path)
        preview = ImportPreview(
            source_name=path.name,
            source_file_hash=self._sha256(path),
            source_version=None,
        )
        archive_error = self._validate_archive(path)
        if archive_error:
            preview.issues.append(ImportIssue("workbook", None, None, archive_error))
            return preview
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            preview.issues.append(ImportIssue("workbook", None, None, f"No fue posible abrir el Excel: {exc}"))
            return preview

        try:
            preview.source_version = self._source_version(workbook)
            missing = {"imp-questions", "imp-answers"} - set(workbook.sheetnames)
            for sheet in sorted(missing):
                preview.issues.append(ImportIssue(sheet, None, None, "La hoja obligatoria no existe."))
            if missing:
                return preview

            preview.answer_sets = self._parse_answer_sets(workbook["imp-answers"], preview.issues)
            answer_sets_by_code = {item["external_code"]: item for item in preview.answer_sets}
            preview.questions = self._parse_questions(
                workbook["imp-questions"], answer_sets_by_code, preview.issues
            )
            return preview
        finally:
            workbook.close()

    @staticmethod
    def _validate_archive(path: Path) -> str | None:
        if not zipfile.is_zipfile(path):
            return "El archivo no es un contenedor XLSX válido."
        try:
            with zipfile.ZipFile(path) as archive:
                members = archive.infolist()
                if len(members) > 1000:
                    return "El archivo XLSX contiene demasiados elementos internos."
                total_uncompressed = sum(member.file_size for member in members)
                if total_uncompressed > 100 * 1024 * 1024:
                    return "El contenido descomprimido del XLSX supera el límite de seguridad."
                for member in members:
                    if member.file_size > 40 * 1024 * 1024:
                        return "El XLSX contiene un elemento interno excesivamente grande."
                    if member.compress_size and member.file_size / member.compress_size > 200:
                        return "El XLSX presenta una relación de compresión no permitida."
        except (OSError, zipfile.BadZipFile) as exc:
            return f"No fue posible inspeccionar el contenedor XLSX: {exc}"
        return None

    @staticmethod
    def _sha256(path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(chunk)
        return digest.hexdigest()

    @staticmethod
    def _source_version(workbook) -> str | None:
        if "Attribution and License" not in workbook.sheetnames:
            return None
        value = workbook["Attribution and License"]["B3"].value
        if value is None:
            return None
        return str(value).strip().lstrip("'").lstrip("vV") or None

    def _parse_answer_sets(self, sheet, issues: list[ImportIssue]) -> list[dict[str, Any]]:
        headers = tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
        if headers[: len(ANSWER_HEADERS)] != ANSWER_HEADERS:
            issues.append(
                ImportIssue(
                    sheet.title,
                    1,
                    None,
                    f"Encabezados inválidos. Se esperaban: {', '.join(ANSWER_HEADERS)}.",
                )
            )
            return []

        answer_sets: list[dict[str, Any]] = []
        seen: set[str] = set()
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(value is not None and str(value).strip() for value in row):
                continue
            code = self._clean_code(row[0])
            if not code:
                issues.append(ImportIssue(sheet.title, row_number, "ANS_SET_CODE", "El código es obligatorio."))
                continue
            if code in seen:
                issues.append(ImportIssue(sheet.title, row_number, "ANS_SET_CODE", f"Código duplicado: {code}."))
                continue
            seen.add(code)

            options: list[dict[str, Any]] = []
            for index, option_code in enumerate(("A", "B", "C", "D")):
                text = self._clean_text(row[index + 1] if len(row) > index + 1 else None)
                raw_weight = row[index + 5] if len(row) > index + 5 else None
                if not text and raw_weight is None:
                    continue
                if not text:
                    issues.append(ImportIssue(sheet.title, row_number, option_code, "La alternativa requiere texto."))
                    continue
                try:
                    weight = Decimal(str(raw_weight))
                except (InvalidOperation, TypeError):
                    issues.append(ImportIssue(sheet.title, row_number, f"{option_code}_W", "La ponderación no es numérica."))
                    continue
                if weight < 0 or weight > 1:
                    issues.append(ImportIssue(sheet.title, row_number, f"{option_code}_W", "La ponderación debe estar entre 0 y 1."))
                options.append(
                    {
                        "option_code": option_code,
                        "text": text,
                        "weight": self._decimal_text(weight),
                        "sort_order": index + 1,
                    }
                )
            if len(options) < 2:
                issues.append(ImportIssue(sheet.title, row_number, None, "El conjunto requiere al menos dos alternativas."))
                continue
            weights = [Decimal(option["weight"]) for option in options]
            if weights != sorted(weights):
                issues.append(ImportIssue(sheet.title, row_number, None, "Las ponderaciones no están ordenadas de menor a mayor.", "warning"))
            content_hash = self._content_hash(options)
            answer_sets.append(
                {
                    "external_code": code,
                    "name": f"SAMM Answer Set {code}",
                    "description": "Conjunto importado desde la hoja imp-answers.",
                    "content_hash": content_hash,
                    "options": options,
                }
            )
        return answer_sets

    def _parse_questions(
        self,
        sheet,
        answer_sets_by_code: dict[str, dict[str, Any]],
        issues: list[ImportIssue],
    ) -> list[dict[str, Any]]:
        headers = tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
        if headers[: len(QUESTION_HEADERS)] != QUESTION_HEADERS:
            issues.append(
                ImportIssue(
                    sheet.title,
                    1,
                    None,
                    f"Encabezados inválidos. Se esperaban: {', '.join(QUESTION_HEADERS)}.",
                )
            )
            return []

        questions: list[dict[str, Any]] = []
        seen: set[str] = set()
        hierarchy_names: dict[str, str] = {}
        function_order: dict[str, int] = {}
        practice_order: dict[str, int] = {}
        stream_order: dict[str, int] = {}

        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(value is not None and str(value).strip() for value in row):
                continue
            values = list(row[: len(QUESTION_HEADERS)])
            code = self._clean_text(values[0]).upper()
            function_name = self._clean_text(values[1])
            practice_name = self._clean_text(values[2])
            stream_name = self._clean_text(values[3])
            question_text = self._clean_text(values[5])
            guidance = self._clean_text(values[6])
            answer_code = self._clean_code(values[7])

            required = {
                "ID": code,
                "Business Function": function_name,
                "Security Practice": practice_name,
                "Activity": stream_name,
                "Question": question_text,
                "Answer Option": answer_code,
            }
            missing = [name for name, value in required.items() if not value]
            if missing:
                for field_name in missing:
                    issues.append(ImportIssue(sheet.title, row_number, field_name, "El campo es obligatorio."))
                continue
            if code in seen:
                issues.append(ImportIssue(sheet.title, row_number, "ID", f"Pregunta duplicada: {code}."))
                continue
            seen.add(code)
            if not QUESTION_ID_PATTERN.match(code):
                issues.append(ImportIssue(sheet.title, row_number, "ID", "El identificador no cumple el formato SAMM esperado.", "warning"))

            try:
                maturity = int(values[4])
            except (TypeError, ValueError):
                issues.append(ImportIssue(sheet.title, row_number, "Maturity", "El nivel de madurez debe ser entero."))
                continue
            if maturity < 1:
                issues.append(ImportIssue(sheet.title, row_number, "Maturity", "El nivel de madurez debe ser mayor que cero."))
                continue
            if answer_code not in answer_sets_by_code:
                issues.append(ImportIssue(sheet.title, row_number, "Answer Option", f"No existe el conjunto de respuestas {answer_code}."))
                continue

            segments = code.split("-")
            function_code = segments[0]
            practice_code = "-".join(segments[:2])
            stream_code = "-".join(segments[:3])
            self._validate_hierarchy_name(hierarchy_names, function_code, function_name, sheet.title, row_number, issues)
            self._validate_hierarchy_name(hierarchy_names, practice_code, practice_name, sheet.title, row_number, issues)
            self._validate_hierarchy_name(hierarchy_names, stream_code, stream_name, sheet.title, row_number, issues)
            function_order.setdefault(function_code, len(function_order) + 1)
            practice_order.setdefault(practice_code, len(practice_order) + 1)
            stream_order.setdefault(stream_code, len(stream_order) + 1)

            criteria = [line.strip(" \t\r\n•-–") for line in guidance.splitlines() if line.strip(" \t\r\n•-–")]
            normalized = {
                "external_code": code,
                "canonical_name": question_text[:240],
                "business_function_code": function_code,
                "business_function_name": function_name,
                "business_function_sort_order": function_order[function_code],
                "security_practice_code": practice_code,
                "security_practice_name": practice_name,
                "security_practice_sort_order": practice_order[practice_code],
                "practice_stream_code": stream_code,
                "practice_stream_name": stream_name,
                "practice_stream_sort_order": stream_order[stream_code],
                "maturity_level": maturity,
                "maturity_name": f"Nivel {maturity}",
                "question_text": question_text,
                "guidance_text": guidance or None,
                "criteria": criteria,
                "answer_set_external_code": answer_code,
                "answer_set_content_hash": answer_sets_by_code[answer_code]["content_hash"],
                "source_row": row_number,
                "sort_order": len(questions) + 1,
            }
            normalized["content_hash"] = self._content_hash(
                {key: value for key, value in normalized.items() if key not in {"source_row", "sort_order"}}
            )
            questions.append(normalized)
        return questions

    @staticmethod
    def _validate_hierarchy_name(
        names: dict[str, str], code: str, name: str, sheet: str, row: int, issues: list[ImportIssue]
    ) -> None:
        previous = names.setdefault(code, name)
        if previous != name:
            issues.append(
                ImportIssue(
                    sheet,
                    row,
                    None,
                    f"El código {code} aparece con nombres distintos: '{previous}' y '{name}'.",
                )
            )

    @staticmethod
    def _clean_text(value: Any) -> str:
        if value is None:
            return ""
        return " ".join(str(value).replace("\u00a0", " ").strip().split()) if "\n" not in str(value) else "\n".join(line.strip() for line in str(value).replace("\u00a0", " ").strip().splitlines())

    @staticmethod
    def _clean_code(value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()

    @staticmethod
    def _decimal_text(value: Decimal) -> str:
        return format(value.normalize(), "f") if value != 0 else "0"

    @staticmethod
    def _content_hash(value: Any) -> str:
        payload = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        return hashlib.sha256(payload.encode("utf-8")).hexdigest()


samm_workbook_parser = SammWorkbookParser()
