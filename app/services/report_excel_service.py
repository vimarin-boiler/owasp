from __future__ import annotations

from io import BytesIO
from typing import Iterable

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


NTT_BLUE = "0072BC"
NTT_CYAN = "00A7E1"
NTT_DARK = "0B1F33"
NTT_LIGHT = "F4F7FA"
WHITE = "FFFFFF"
SUCCESS = "198754"
WARNING = "FFC107"
DANGER = "DC3545"
MUTED = "6C757D"


class ReportExcelService:
    def build(self, payload: dict) -> BytesIO:
        workbook = Workbook()
        workbook.remove(workbook.active)
        self._summary(workbook, payload)
        self._dimensions(workbook, "Funciones", payload["functions"])
        self._dimensions(workbook, "Prácticas", payload["practices"])
        self._dimensions(workbook, "Flujos", payload["streams"])
        self._questions(workbook, payload["questions"])
        self._evidences(workbook, payload["evidences"])
        self._recommendations(workbook, payload["recommendations"])
        self._roadmap(workbook, payload["roadmap"])
        self._reviews(workbook, payload["reviews"])
        self._history(workbook, payload["response_history"])
        self._audit(workbook, payload["audit_log"])
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return output

    def _new_sheet(self, workbook: Workbook, title: str):
        sheet = workbook.create_sheet(title=title[:31])
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A2"
        return sheet

    @staticmethod
    def _title(sheet, title: str, subtitle: str = "") -> int:
        sheet.merge_cells("A1:H1")
        cell = sheet["A1"]
        cell.value = title
        cell.font = Font(name="Aptos Display", size=20, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NTT_DARK)
        cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 34
        if subtitle:
            sheet.merge_cells("A2:H2")
            sheet["A2"] = subtitle
            sheet["A2"].font = Font(name="Aptos", size=10, color=MUTED)
            sheet["A2"].alignment = Alignment(wrap_text=True)
            return 4
        return 3

    @staticmethod
    def _header(sheet, row: int, headers: Iterable[str]) -> None:
        for col, header in enumerate(headers, 1):
            cell = sheet.cell(row=row, column=col, value=header)
            cell.font = Font(name="Aptos", bold=True, color=WHITE)
            cell.fill = PatternFill("solid", fgColor=NTT_BLUE)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[row].height = 26

    @staticmethod
    def _finish(sheet, header_row: int, last_row: int, last_col: int, table_name: str | None = None) -> None:
        thin = Side(style="thin", color="DDE5EC")
        for row in sheet.iter_rows(min_row=header_row + 1, max_row=max(last_row, header_row), min_col=1, max_col=last_col):
            for cell in row:
                cell.font = Font(name="Aptos", size=10, color=NTT_DARK)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = Border(bottom=thin)
        for column in range(1, last_col + 1):
            max_length = 12
            for cell in sheet[get_column_letter(column)]:
                if cell.value is not None:
                    max_length = max(max_length, min(len(str(cell.value)) + 2, 58))
            sheet.column_dimensions[get_column_letter(column)].width = max_length
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(last_col)}{max(last_row, header_row)}"
        sheet.freeze_panes = f"A{header_row + 1}"
        if table_name and last_row > header_row:
            table = Table(displayName=table_name, ref=f"A{header_row}:{get_column_letter(last_col)}{last_row}")
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
                showRowStripes=True, showColumnStripes=False,
            )
            sheet.add_table(table)

    def _summary(self, workbook: Workbook, payload: dict) -> None:
        sheet = self._new_sheet(workbook, "Resumen ejecutivo")
        assessment = payload["assessment"]
        result = payload["result"]
        sheet.merge_cells("A1:H2")
        company = payload.get("branding", {}).get("company_name", "NTT DATA")
        sheet["A1"] = f"{company.upper()} DEVSECOPS ASSESSMENT"
        sheet["A1"].font = Font(name="Aptos Display", size=22, bold=True, color=WHITE)
        sheet["A1"].fill = PatternFill("solid", fgColor=NTT_DARK)
        sheet["A1"].alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 38
        sheet.row_dimensions[2].height = 16
        sheet["A4"] = assessment["organization"]
        sheet["A4"].font = Font(name="Aptos Display", size=18, bold=True, color=NTT_BLUE)
        sheet.merge_cells("A4:H4")
        sheet["A5"] = assessment["name"]
        sheet["A5"].font = Font(name="Aptos", size=13, bold=True, color=NTT_DARK)
        sheet.merge_cells("A5:H5")
        metrics = [
            ("Madurez actual", result["overall_score"]),
            ("Nivel objetivo", result["target_score"]),
            ("Brecha", result["gap"]),
            ("Avance", f"{result['progress_percent']:.1f}%"),
        ]
        for index, (label, value) in enumerate(metrics):
            col = 1 + index * 2
            sheet.merge_cells(start_row=7, start_column=col, end_row=7, end_column=col + 1)
            sheet.merge_cells(start_row=8, start_column=col, end_row=9, end_column=col + 1)
            sheet.cell(7, col, label).font = Font(name="Aptos", size=9, bold=True, color=MUTED)
            sheet.cell(8, col, "N/D" if value is None else value).font = Font(name="Aptos Display", size=20, bold=True, color=NTT_BLUE)
            sheet.cell(8, col).alignment = Alignment(vertical="center")
            for row in range(7, 10):
                for column in range(col, col + 2):
                    sheet.cell(row, column).fill = PatternFill("solid", fgColor=NTT_LIGHT)
        details = [
            ("Estado", assessment["status"]),
            ("Versión SAMM", assessment["questionnaire_version"]),
            ("Fuente de cálculo", assessment["scoring_source"]),
            ("Fecha de inicio", assessment["start_date"] or ""),
            ("Fecha objetivo", assessment["target_date"] or ""),
            ("Snapshot", result["snapshot_number"] or "Vista previa"),
            ("Hash de cálculo", result["input_hash"]),
            ("Generado", payload["generated_at"]),
        ]
        row = 12
        for label, value in details:
            sheet.cell(row, 1, label).font = Font(bold=True, color=MUTED)
            sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
            sheet.cell(row, 2, value)
            row += 1
        self._header(sheet, row + 1, ["Función", "Actual", "Objetivo", "Brecha", "Avance %", "Aplicables", "No aplica", "Pendientes"])
        header_row = row + 1
        for item in payload["functions"]:
            row += 1
            values = [item["name"], item["score"], item["target_score"], item["gap"], item["normalized_percent"], item["applicable_count"], item["not_applicable_count"], item["pending_count"]]
            for col, value in enumerate(values, 1):
                sheet.cell(row + 1, col, value)
        last_row = row + 1
        if last_row > header_row:
            sheet.conditional_formatting.add(
                f"D{header_row + 1}:D{last_row}",
                ColorScaleRule(start_type="min", start_color=SUCCESS, mid_type="percentile", mid_value=50, mid_color=WARNING, end_type="max", end_color=DANGER),
            )
        self._finish(sheet, header_row, last_row, 8, "ExecutiveFunctions")

    def _dimensions(self, workbook: Workbook, title: str, items: list[dict]) -> None:
        sheet = self._new_sheet(workbook, title)
        row = self._title(sheet, f"Resultados por {title.lower()}")
        headers = ["Código", "Nombre", "Función", "Práctica", "Actual", "Objetivo", "Brecha", "Avance %", "Aplicables", "No aplica", "Pendientes"]
        self._header(sheet, row, headers)
        for item in items:
            metadata = item.get("metadata", {})
            row += 1
            values = [
                metadata.get("stream_code") or metadata.get("practice_code") or metadata.get("function_code") or item["key"],
                item["name"], metadata.get("function_name", ""), metadata.get("practice_name", ""),
                item["score"], item["target_score"], item["gap"], item["normalized_percent"],
                item["applicable_count"], item["not_applicable_count"], item["pending_count"],
            ]
            for col, value in enumerate(values, 1):
                sheet.cell(row, col, value)
        self._finish(sheet, 3, row, len(headers), f"Table{title.replace('á','a').replace('ó','o')}")
        if row > 3:
            sheet.conditional_formatting.add(
                f"G4:G{row}",
                ColorScaleRule(start_type="min", start_color=SUCCESS, mid_type="percentile", mid_value=50, mid_color=WARNING, end_type="max", end_color=DANGER),
            )

    def _questions(self, workbook: Workbook, rows: list[dict]) -> None:
        sheet = self._new_sheet(workbook, "Preguntas")
        header_row = self._title(sheet, "Detalle de preguntas y respuestas")
        headers = ["Código", "Función", "Práctica", "Flujo", "Nivel", "Pregunta", "Estado", "Respuesta", "Ponderación", "No aplica", "Justificación N/A", "Comentario respondedor", "Comentario revisor", "Respondedor", "Revisor", "Evidencias"]
        self._header(sheet, header_row, headers)
        row = header_row
        for item in rows:
            row += 1
            values = [item["code"], item["function"], item["practice"], item["stream"], item["level"], item["question"], item["status_label"], item["answer"], item["weight"], "Sí" if item["not_applicable"] else "No", item["not_applicable_justification"], item["respondent_comment"], item["reviewer_comment"], item["respondent"], item["reviewer"], item["evidence_count"]]
            for col, value in enumerate(values, 1):
                sheet.cell(row, col, value)
        self._finish(sheet, header_row, row, len(headers), "AssessmentQuestions")

    def _evidences(self, workbook: Workbook, rows: list[dict]) -> None:
        sheet = self._new_sheet(workbook, "Evidencias")
        header_row = self._title(sheet, "Inventario de evidencias")
        headers = ["Pregunta", "Archivo", "Descripción", "Extensión", "MIME", "Tamaño bytes", "SHA-256", "Cargado por", "Fecha", "Validación", "Comentario de revisión", "Revisado por"]
        self._header(sheet, header_row, headers)
        row = header_row
        for item in rows:
            row += 1
            values = [item["question_code"], item["filename"], item["description"], item["extension"], item["mime_type"], item["size_bytes"], item["sha256"], item["uploaded_by"], item["uploaded_at"], item["validation_status"], item["review_comment"], item["reviewed_by"]]
            for col, value in enumerate(values, 1): sheet.cell(row, col, value)
        self._finish(sheet, header_row, row, len(headers), "AssessmentEvidence")

    def _recommendations(self, workbook: Workbook, rows: list[dict]) -> None:
        sheet = self._new_sheet(workbook, "Recomendaciones")
        header_row = self._title(sheet, "Recomendaciones priorizadas")
        headers = ["Prioridad", "Título", "Descripción", "Riesgo", "Esfuerzo", "Responsable", "Horizonte", "Fecha objetivo", "Dependencias", "Estado", "Quick win", "Nivel objetivo"]
        self._header(sheet, header_row, headers)
        row = header_row
        for item in rows:
            row += 1
            values = [item["priority_label"], item["title"], item["description"], item["risk"], item["effort"], item["owner"], item["horizon"], item["due_date"], item["dependencies"], item["status_label"], "Sí" if item["quick_win"] else "No", item["target_level"]]
            for col, value in enumerate(values, 1): sheet.cell(row, col, value)
        self._finish(sheet, header_row, row, len(headers), "AssessmentRecommendations")

    def _roadmap(self, workbook: Workbook, groups: list[dict]) -> None:
        sheet = self._new_sheet(workbook, "Roadmap")
        header_row = self._title(sheet, "Roadmap de mejora")
        headers = ["Horizonte", "Orden", "Prioridad", "Iniciativa", "Responsable", "Esfuerzo", "Estado", "Quick win", "Fecha objetivo"]
        self._header(sheet, header_row, headers)
        row = header_row
        for group_order, group in enumerate(groups, 1):
            for item_order, item in enumerate(group["items"], 1):
                row += 1
                values = [group["horizon"], f"{group_order}.{item_order}", item["priority_label"], item["title"], item["owner"], item["effort"], item["status_label"], "Sí" if item["quick_win"] else "No", item["due_date"]]
                for col, value in enumerate(values, 1): sheet.cell(row, col, value)
        self._finish(sheet, header_row, row, len(headers), "AssessmentRoadmap")

    def _reviews(self, workbook: Workbook, rows: list[dict]) -> None:
        sheet = self._new_sheet(workbook, "Revisiones")
        header_row = self._title(sheet, "Historial de revisiones")
        headers = ["Pregunta", "Versión respuesta", "Decisión", "Revisor", "Fecha", "Comentario"]
        self._header(sheet, header_row, headers)
        row = header_row
        for item in rows:
            row += 1
            values = [item["question_code"], item["response_version"], item["decision"], item["reviewer"], item["reviewed_at"], item["comment"]]
            for col, value in enumerate(values, 1): sheet.cell(row, col, value)
        self._finish(sheet, header_row, row, len(headers), "AssessmentReviews")

    def _history(self, workbook: Workbook, rows: list[dict]) -> None:
        sheet = self._new_sheet(workbook, "Historial respuestas")
        header_row = self._title(sheet, "Trazabilidad de respuestas")
        headers = ["Pregunta", "Versión", "Desde", "Hacia", "Usuario", "Fecha", "Motivo"]
        self._header(sheet, header_row, headers)
        row = header_row
        for item in rows:
            row += 1
            values = [item["question_code"], item["version"], item["from"], item["to"], item["changed_by"], item["changed_at"], item["reason"]]
            for col, value in enumerate(values, 1): sheet.cell(row, col, value)
        self._finish(sheet, header_row, row, len(headers), "ResponseHistory")

    def _audit(self, workbook: Workbook, rows: list[dict]) -> None:
        sheet = self._new_sheet(workbook, "Bitácora")
        header_row = self._title(sheet, "Bitácora del assessment")
        headers = ["Fecha", "Usuario", "Acción", "Entidad", "ID público", "Resultado", "IP"]
        self._header(sheet, header_row, headers)
        row = header_row
        for item in rows:
            row += 1
            values = [item["occurred_at"], item["actor"], item["action"], item["entity_type"], item["entity_public_id"], item["result"], item["ip_address"]]
            for col, value in enumerate(values, 1): sheet.cell(row, col, value)
        self._finish(sheet, header_row, row, len(headers), "AssessmentAudit")


report_excel_service = ReportExcelService()
