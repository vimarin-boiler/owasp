from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import selectinload

from app.extensions import db
from app.models import QuestionnaireVersion, QuestionnaireVersionQuestion, QuestionRevision


class CatalogExportService:
    def build_workbook(self, version: QuestionnaireVersion) -> BytesIO:
        loaded = db.session.scalar(
            select(QuestionnaireVersion)
            .options(
                selectinload(QuestionnaireVersion.question_links)
                .selectinload(QuestionnaireVersionQuestion.question_revision)
                .selectinload(QuestionRevision.criteria),
                selectinload(QuestionnaireVersion.question_links)
                .selectinload(QuestionnaireVersionQuestion.question_revision)
                .selectinload(QuestionRevision.answer_set),
                selectinload(QuestionnaireVersion.question_links)
                .selectinload(QuestionnaireVersionQuestion.question_revision)
                .selectinload(QuestionRevision.business_function),
                selectinload(QuestionnaireVersion.question_links)
                .selectinload(QuestionnaireVersionQuestion.question_revision)
                .selectinload(QuestionRevision.security_practice),
                selectinload(QuestionnaireVersion.question_links)
                .selectinload(QuestionnaireVersionQuestion.question_revision)
                .selectinload(QuestionRevision.practice_stream),
                selectinload(QuestionnaireVersion.question_links)
                .selectinload(QuestionnaireVersionQuestion.question_revision)
                .selectinload(QuestionRevision.maturity_level),
            )
            .where(QuestionnaireVersion.id == version.id)
        )
        if loaded is None:
            raise ValueError("La versión no existe.")

        workbook = Workbook()
        metadata = workbook.active
        metadata.title = "Metadata"
        metadata.append(["Campo", "Valor"])
        metadata.append(["Nombre", loaded.name])
        metadata.append(["Versión", loaded.version_number])
        metadata.append(["Estado", loaded.status.value])
        metadata.append(["Descripción", loaded.description or ""])
        metadata.append(["Fuente", loaded.source_name or ""])
        metadata.append(["Hash SHA-256", loaded.source_file_hash or ""])

        questions_sheet = workbook.create_sheet("imp-questions")
        questions_sheet.append(["ID", "Business Function", "Security Practice", "Activity", "Maturity", "Question", "Guidance", "Answer Option"])
        answer_sets: dict[int, tuple[int, object]] = {}
        next_answer_code = 0
        for link in sorted(loaded.question_links, key=lambda item: item.sort_order):
            revision = link.question_revision
            if revision.answer_set_id not in answer_sets:
                answer_sets[revision.answer_set_id] = (next_answer_code, revision.answer_set)
                next_answer_code += 1
            answer_code = answer_sets[revision.answer_set_id][0]
            questions_sheet.append([
                revision.question.external_code if revision.question else "",
                revision.business_function.name,
                revision.security_practice.name,
                revision.practice_stream.name,
                revision.maturity_level.level_number,
                revision.question_text,
                "\n".join(criterion.criterion_text for criterion in revision.criteria) or revision.guidance_text or "",
                answer_code,
            ])

        answers_sheet = workbook.create_sheet("imp-answers")
        answers_sheet.append(["ANS_SET_CODE", "A", "B", "C", "D", "A_W", "B_W", "C_W", "D_W"])
        for _, (answer_code, answer_set) in sorted(answer_sets.items(), key=lambda item: item[1][0]):
            options = sorted(answer_set.options, key=lambda option: option.sort_order)
            texts = [option.text for option in options[:4]]
            weights = [float(option.weight) for option in options[:4]]
            answers_sheet.append([answer_code, *texts, *([None] * (4 - len(texts))), *weights, *([None] * (4 - len(weights)))])

        for sheet in (metadata, questions_sheet, answers_sheet):
            header = sheet[1]
            for cell in header:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="0072BC")
                cell.alignment = Alignment(horizontal="center")
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for column in sheet.columns:
                letter = column[0].column_letter
                width = min(max(len(str(cell.value or "")) for cell in column) + 2, 55)
                sheet.column_dimensions[letter].width = max(12, width)
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)

        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return output


catalog_export_service = CatalogExportService()
